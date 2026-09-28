"""Lo que la gente busca en la página y no encuentra (en el buscador del catálogo y en el chat
del muñequito), con lo más pedido primero. Sirve para decidir qué mercadería traer.

La página anota en Umami cada búsqueda sin resultados (eventos "busqueda-sin-resultados" y
"chat-sin-resultados", con lo que se escribió en la propiedad "q"). Este script los junta y
arma un Excel: busquedas.xlsx.

Necesita en el archivo .env (una sola vez):
    UMAMI_API_KEY=...      (en Umami: Settings -> API keys -> Create key)

Solo lee datos. El Excel se queda en esta carpeta (está en .gitignore).
Sin clave también se puede ver en Umami: Eventos -> busqueda-sin-resultados -> propiedad "q".
"""

import os
import sys
import time
import unicodedata
from collections import Counter, defaultdict

import requests
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Font

load_dotenv()
API_KEY = os.environ.get("UMAMI_API_KEY", "").strip()
SITIO = "1f84c623-7a7b-4d1f-90cd-ef1e1a646852"   # el mismo de CONFIG.analitica en docs/index.html
BASE = "https://api.umami.is/v1"
DIAS = int(sys.argv[1]) if len(sys.argv) > 1 else 30
EVENTOS = {"busqueda-sin-resultados": "Buscador del catálogo", "chat-sin-resultados": "Chat del muñequito"}
SALIDA = "busquedas.xlsx"


def valores(evento, desde, hasta):
    """[(lo que escribieron, cuántas veces)] de un evento, entre dos fechas (milisegundos)."""
    r = requests.get(
        f"{BASE}/websites/{SITIO}/event-data/values",
        headers={"x-umami-api-key": API_KEY, "Accept": "application/json"},
        params={"startAt": desde, "endAt": hasta, "event": evento, "eventName": evento, "propertyName": "q"},
        timeout=60,
    )
    if r.status_code in (401, 403):
        sys.exit("Umami no aceptó la clave. Revisa UMAMI_API_KEY en el archivo .env.")
    if r.status_code != 200:
        sys.exit(f"Umami respondió {r.status_code}: {r.text[:200]}")
    datos = r.json()
    filas = datos if isinstance(datos, list) else datos.get("data", [])
    return [(str(f.get("value", "")).strip(), int(f.get("total", 0) or 0)) for f in filas]


def clave(texto):
    """'Zelda  Tears' y 'zelda tears' cuentan como lo mismo."""
    t = unicodedata.normalize("NFD", texto.lower())
    return " ".join("".join(c for c in t if unicodedata.category(c) != "Mn").split())


def main():
    if not API_KEY:
        sys.exit("Falta UMAMI_API_KEY en el archivo .env (Umami: Settings -> API keys -> Create key).")
    hasta = int(time.time() * 1000)
    desde = hasta - DIAS * 86400000
    veces, donde, como = Counter(), defaultdict(set), {}
    for evento, lugar in EVENTOS.items():
        for texto, n in valores(evento, desde, hasta):
            k = clave(texto)
            if not k:
                continue
            veces[k] += n
            donde[k].add(lugar)
            como.setdefault(k, texto)

    wb = Workbook()
    ws = wb.active
    ws.title = "Lo que buscan"
    ws.append([f"Lo que buscaron en la página y no encontraron (últimos {DIAS} días)"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(["Lo que buscaron", "Veces", "Dónde"])
    for c in ws[3]:
        c.font = Font(bold=True)
    for k, n in veces.most_common():
        ws.append([como[k], n, ", ".join(sorted(donde[k]))])
    ws.column_dimensions["A"].width = 45
    ws.column_dimensions["B"].width = 10
    ws.column_dimensions["C"].width = 40
    wb.save(SALIDA)
    print(f"{len(veces)} búsquedas distintas sin resultados. Guardado en {SALIDA}.")


if __name__ == "__main__":
    main()
